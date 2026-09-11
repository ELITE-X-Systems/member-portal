"""
Generate the public ELITE X portal data from the PRIVATE Staff Records workbook.

Usage:
  python scripts/build_data.py "Elite-X Records.xlsx"

The generated data.json contains only member-facing fields.
Review it before publishing.
"""
import sys, json
from datetime import datetime
from openpyxl import load_workbook

if len(sys.argv) != 2:
    raise SystemExit("Usage: python scripts/build_data.py <private-workbook.xlsx>")

src = sys.argv[1]
wb = load_workbook(src, data_only=True)
reg, bond, ded, arch = (wb[x] for x in ("Member Registry","Bond Record","Deductions","Monthly Archive"))

members = {}
for r in range(6,106):
    mid, name = reg.cell(r,1).value, reg.cell(r,2).value
    if not (mid and name):
        continue
    br = r + 3
    entries = []
    for c, week, entry in [(4,"Week 1",1),(5,"Week 1",2),(6,"Week 1",3),(7,"Week 1",4),
                           (8,"Week 2",1),(9,"Week 2",2),(10,"Week 2",3),(11,"Week 2",4)]:
        value = bond.cell(br,c).value
        if value not in (None,0,""):
            entries.append({"week":week,"entry":entry,"bonds":float(value)})
    members[str(mid)] = {
        "id":str(mid), "name":str(name),
        "rank":reg.cell(r,3).value or "Initiate",
        "idStatus":reg.cell(r,4).value or "",
        "monthlyEarned":float(bond.cell(br,12).value or 0),
        "deductions":float(bond.cell(br,13).value or 0),
        "overallNet":float(bond.cell(br,14).value or 0),
        "status":bond.cell(br,15).value or "",
        "weeklyEntries":entries
    }

transactions=[]
for r in range(7,207):
    mid, date, typ, desc, amount = ded.cell(r,3).value, ded.cell(r,2).value, ded.cell(r,5).value, ded.cell(r,6).value, ded.cell(r,7).value
    if mid and typ and amount is not None:
        transactions.append({
            "memberId":str(mid),
            "date":date.strftime("%Y-%m-%d") if isinstance(date,datetime) else str(date or ""),
            "type":str(typ).strip(),
            "description":str(desc or "").strip(),
            "amount":float(amount)
        })

monthly={}
for r in range(6,106):
    mid=arch.cell(r,1).value
    if mid and str(mid) in members:
        monthly[str(mid)] = {
            "September":float(arch.cell(r,3).value or 0),
            "October":float(arch.cell(r,4).value or 0),
            "November":float(arch.cell(r,5).value or 0),
            "December":float(arch.cell(r,6).value or 0)
        }

out={
    "version":"1.0",
    "currency":"BONDS 💴",
    "source":"ELITE X Private Staff Records",
    "generatedAt":datetime.now().strftime("%Y-%m-%d"),
    "members":list(members.values()),
    "transactions":transactions,
    "monthlyArchive":monthly,
    "activityLog":[],
    "activityMappingNote":"Current Activity Log has no Member ID, so detailed activity-log rows are not attributed to members."
}
with open("data.json","w",encoding="utf-8") as f:
    json.dump(out,f,ensure_ascii=False,indent=2)
print("Wrote sanitized data.json")
